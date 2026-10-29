"""Build sicav.json — the SICAV market from a depositary point of view.

Sources
  * CNMV SOCREGISTRO / SOCTRIM / SOCPATRIMDISVAR (quarterly XML, per SICAV):
      depositario, gestora, AUM, shareholders, TER, returns and the depositary /
      management fees each SICAV actually paid (see parsers/cnmv_sicav.py).
  * CNMV Estadísticas IIC file 3: Cuadro 4.1 (quarterly SICAV totals) and Cuadro 8.4
      (sector P&L — total depositary and management fees paid by all SICAVs).
  * Inverco monthly `*_00-SICAV.xlsx`: monthly SICAV totals and group ranking
      (includes sociedades de inversión libre; ~99 % sample) — fresher than CNMV.
"""
import glob
import json
import os
import re

import openpyxl

from ..config import DATA_DIR, OUTPUT_DIR, get_snapshot_folders
from ..parsers.cnmv_sicav import parse_cnmv_sicav_full
from .cnmv_depositaria import _is_captive

CNMV_DIR = os.path.join(DATA_DIR, 'CNMV Estadisticas')
INVERSIS = 'BANCO INVERSIS, S.A.'


_LEGAL_TOKENS = re.compile(
    r'(?<![^\s,])(SICAV|S\.I\.C\.A\.V\.|S\.A\.U\.|SAU|S\.A\.|S\.A|SA|SGIIC|S\.G\.I\.I\.C\.)(?![^\s,])',
    re.I)


def _short(name):
    """Drop legal-form tokens (SICAV, S.A., SGIIC …) — whole tokens only."""
    s = _LEGAL_TOKENS.sub('', name or '')
    s = re.sub(r'\s*,\s*(,\s*)*', ', ', s)
    return re.sub(r'\s+', ' ', s).strip(' ,')


def _wavg(pairs):
    """AUM-weighted average of (weight, value) pairs, skipping None values."""
    num = sum(w * v for w, v in pairs if v is not None)
    den = sum(w for w, v in pairs if v is not None)
    return num / den if den > 0 else None


def _median(vals):
    s = sorted(v for v in vals if v is not None)
    if not s:
        return None
    n = len(s)
    return s[n // 2] if n % 2 else (s[n // 2 - 1] + s[n // 2]) / 2


def _bps(pct):
    return round(pct * 100, 2) if pct is not None else None


# ── CNMV Estadísticas file 3 ──────────────────────────────────────────────

_ROMAN = {'I': 1, 'II': 2, 'III': 3, 'IV': 4}


def _cnmv_file3_paths():
    pat = re.compile(r'^Estadisticas_IIC_(\d{4})_(\d)T_3\.xlsx$')
    out = []
    for name in os.listdir(CNMV_DIR):
        m = pat.match(name)
        if m:
            out.append(((int(m.group(1)), int(m.group(2))), os.path.join(CNMV_DIR, name)))
    return sorted(out)


def _parse_cuadro_41(path):
    """Quarterly SICAV count / accounts / AUM from Cuadro 4.1 (5 quarters)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = [r for r in wb['Cuadro 4.1'].iter_rows(values_only=True)]
    wb.close()
    years_row = quarters_row = None
    for i, r in enumerate(rows):
        if any(str(v).strip() in _ROMAN for v in r if v is not None):
            quarters_row, years_row = r, rows[i - 1]
            break
    if quarters_row is None:
        return {}
    labels, year = {}, None
    for c, q in enumerate(quarters_row):
        y = years_row[c] if c < len(years_row) else None
        if isinstance(y, (int, float)):
            year = int(y)
        if q is not None and str(q).strip() in _ROMAN and year:
            labels[c] = f'{year}-Q{_ROMAN[str(q).strip()]}'
    out = {}
    for r in rows:
        first = str(r[0] or '')
        key = ('count' if first.startswith('1.') and 'mero' in first else
               'accounts' if first.startswith('2.') else
               'aum_k' if first.startswith('3.') else None)
        if not key:
            continue
        for c, lab in labels.items():
            if isinstance(r[c], (int, float)):
                out.setdefault(lab, {})[key] = r[c]
    return out


def _parse_pl_fees(path, sheet):
    """{year: {'dep_k': €k, 'mgmt_k': €k}} from a consolidated P&L cuadro (8.2 FI / 8.4 SICAV)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = [r for r in wb[sheet].iter_rows(values_only=True)]
    wb.close()
    years = None
    out = {}
    for r in rows:
        vals = [v for v in r if v is not None]
        if not vals:
            continue
        if years is None and all(isinstance(v, (int, float)) and 2000 < v < 2100 for v in vals):
            years = [int(v) for v in vals]
            continue
        label = str(vals[0])
        key = 'dep_k' if 'depositario' in label.lower() else 'mgmt_k' if 'n de gesti' in label.lower() else None
        if key and years:
            nums = [v for v in vals[1:] if isinstance(v, (int, float))]
            for y, v in zip(years, nums):
                out.setdefault(y, {})[key] = abs(v)
    return out


# ── Inverco monthly SICAV file ───────────────────────────────────────────

def _parse_inverco_sicav(path):
    """Return (totals, ranking) from an Inverco `*_00-SICAV.xlsx`."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = [r for r in wb.worksheets[0].iter_rows(values_only=True)]
    wb.close()
    totals, ranking, current = None, [], None
    for r in rows:
        r = list(r[:6]) + [None] * (6 - len(r[:6]))
        rank, grupo, gestora, aum, num, holders = r
        label = str(grupo or '').upper().replace(' ', '')
        if 'TOTAL' in label and isinstance(aum, (int, float)):
            totals = {'aum_k': aum, 'count': int(num or 0), 'shareholders': int(holders or 0)}
            break
        if isinstance(rank, (int, float)) and isinstance(aum, (int, float)):
            current = {'rank': int(rank), 'group': str(grupo).strip(), 'gestoras': [],
                       'aum_m': round(aum / 1000, 1), 'count': int(num or 0),
                       'shareholders': int(holders or 0)}
            ranking.append(current)
            if gestora:
                current['gestoras'].append(str(gestora).strip())
        elif current and gestora and isinstance(aum, (int, float)):
            current['gestoras'].append(str(gestora).strip())
    return totals, ranking


def _inverco_monthly():
    series, latest_ranking, latest_month = [], [], None
    for label, folder in get_snapshot_folders():
        files = glob.glob(os.path.join(folder, 'Datos Generales', '*-SICAV.xlsx'))
        if not files:
            continue
        totals, ranking = _parse_inverco_sicav(files[0])
        if not totals:
            continue
        series.append({'month': label, 'aum_bn': round(totals['aum_k'] / 1e6, 2),
                       'count': totals['count'], 'shareholders': totals['shareholders']})
        latest_ranking, latest_month = ranking, label
    return series, latest_ranking, latest_month


# ── Main builder ─────────────────────────────────────────────────────────

def build_sicav_market():
    full = parse_cnmv_sicav_full(CNMV_DIR)
    if not full:
        return {}
    recs = full['records']

    # FI context from cnmv_fees.json (written earlier in the pipeline)
    fees_path = os.path.join(OUTPUT_DIR, 'cnmv_fees.json')
    fi_classes = []
    if os.path.exists(fees_path):
        with open(fees_path, encoding='utf-8') as f:
            fi_classes = [c for c in json.load(f)['funds'] if c.get('fund_type', 'FI') == 'FI']
    fi_by_gestora = {}
    for c in fi_classes:
        g = fi_by_gestora.setdefault(c['gestora'], {'aum_m': 0.0, 'depos': {}, 'grupo': c['grupo']})
        g['aum_m'] += c['patrimonio_m']
        d = c.get('depositario') or ''
        g['depos'][d] = g['depos'].get(d, 0.0) + c['patrimonio_m']
    fi_dep_wtd = _wavg([(c['patrimonio_m'], c['depositary_fee']) for c in fi_classes])
    fi_dep_median = _median([c['depositary_fee'] for c in fi_classes])

    # ── Entity-level (one row per SICAV; series/compartments aggregated) ──
    ents = {}
    for r in recs:
        e = ents.setdefault(r['registro'], {
            'registro': r['registro'], 'name': r['fund_name'], 'short': _short(r['fund_name']),
            'gestora': r['gestora'] or '(Self-managed SICAV)',
            'grupo': (fi_by_gestora.get(r['gestora']) or {}).get('grupo') or r['grupo'],
            'grupo_gestora_cnmv': r['grupo_gestora'],
            'grupo_depositario_cnmv': r['grupo_depositario'],
            'depositario': r['depositario'], 'vocacion': r['vocacion'],
            'aum_k': 0.0, 'investors': 0, 'series': 0, '_w': [],
        })
        e['aum_k'] += r['patrimonio_k']
        e['investors'] += r['investors']
        e['series'] += 1
        e['_w'].append(r)

    prior = full['prior_register']
    prior_aum = full['prior_aum_eur']
    sicavs = []
    for reg, e in ents.items():
        w = e.pop('_w')
        aum_k = e['aum_k']
        weights = [(max(x['patrimonio_k'], 1e-9), x) for x in w]
        dep = _wavg([(wt, x['depositary_fee']) for wt, x in weights])
        p = prior.get(reg)
        sicav = {
            **{k: v for k, v in e.items() if k != 'aum_k'},
            'aum_m': round(aum_k / 1000, 2),
            'dep_fee_bps': _bps(dep),
            'dep_fee_basis': w[0]['dep_fee_basis'],
            'dep_fee_fy_bps': _bps(_wavg([(wt, x['dep_fee_fy']) for wt, x in weights])),
            'dep_fee_ytd_ann_bps': _bps(_wavg([(wt, x['dep_fee_ytd_ann']) for wt, x in weights])),
            'mgmt_fee_pct': round(_wavg([(wt, x['mgmt_fee_aum']) for wt, x in weights]) or 0, 3)
                            if any(x['mgmt_fee_aum'] is not None for x in w) else None,
            'ter_pct': round(_wavg([(wt, x['ter']) for wt, x in weights]), 3)
                       if any(x['ter'] is not None for x in w) else None,
            'return_ytd': _wavg([(wt, x['return_ytd']) for wt, x in weights]),
            'est_dep_fee_k': round(aum_k * dep / 100, 1) if dep is not None else None,
            'is_inversis': e['depositario'] == INVERSIS,
            'is_captive': _is_captive(e['grupo_gestora_cnmv'] or e['grupo'], e['depositario'])
                          or (bool(e['grupo_gestora_cnmv'])
                              and e['grupo_gestora_cnmv'] == e['grupo_depositario_cnmv']),
            'prior_depositario': p['depositario'] if p else None,
            'switched': bool(p) and p['depositario'] != e['depositario'],
            'is_new': p is None and bool(prior),
        }
        if sicav['return_ytd'] is not None:
            sicav['return_ytd'] = round(sicav['return_ytd'], 2)
        sicavs.append(sicav)
    sicavs.sort(key=lambda s: -s['aum_m'])

    total_aum_m = sum(s['aum_m'] for s in sicavs)
    total_fees_k = sum(s['est_dep_fee_k'] or 0 for s in sicavs)

    # ── By depositario ──
    by_depo = {}
    for s in sicavs:
        d = by_depo.setdefault(s['depositario'], {'list': []})
        d['list'].append(s)
    switches = [s for s in sicavs if s['switched']]
    exits = []
    if prior:
        for reg, p in prior.items():
            if reg not in ents and p['depositario']:
                exits.append({
                    'registro': reg, 'name': p['name'], 'short': _short(p['name']),
                    'gestora': p['gestora'], 'depositario': p['depositario'],
                    'aum_m': round(prior_aum.get(reg, 0) / 1e6, 2),
                })
        exits.sort(key=lambda x: -x['aum_m'])
    entries = [s for s in sicavs if s['is_new']]

    depositario_stats = []
    for depo, d in by_depo.items():
        lst = d['list']
        aum = sum(s['aum_m'] for s in lst)
        fees_k = sum(s['est_dep_fee_k'] or 0 for s in lst)
        prior_count = sum(1 for p in prior.values() if p['depositario'] == depo) if prior else None
        depositario_stats.append({
            'depositario': depo, 'short': _short(depo),
            'count': len(lst), 'aum_bn': round(aum / 1000, 3),
            'share_pct': round(aum / total_aum_m * 100, 2) if total_aum_m else 0,
            'count_share_pct': round(len(lst) / len(sicavs) * 100, 1),
            'investors': sum(s['investors'] for s in lst),
            'gestoras': len({s['gestora'] for s in lst}),
            'wtd_fee_bps': _bps(_wavg([(s['aum_m'], (s['dep_fee_bps'] or 0) / 100 if s['dep_fee_bps'] is not None else None) for s in lst])),
            'median_fee_bps': _median([s['dep_fee_bps'] for s in lst]),
            'est_fees_k': round(fees_k, 0),
            'avg_size_m': round(aum / len(lst), 1),
            'prior_count': prior_count,
            'switch_in': sum(1 for s in switches if s['depositario'] == depo),
            'switch_out': sum(1 for s in switches if s['prior_depositario'] == depo),
            'new': sum(1 for s in entries if s['depositario'] == depo),
            'exits': sum(1 for x in exits if x['depositario'] == depo),
            'is_inversis': depo == INVERSIS,
        })
    depositario_stats.sort(key=lambda x: -x['aum_bn'])
    for i, d in enumerate(depositario_stats):
        d['rank'] = i + 1

    # ── By gestora ──
    by_g = {}
    for s in sicavs:
        by_g.setdefault(s['gestora'], []).append(s)
    gestora_stats = []
    for g, lst in by_g.items():
        aum = sum(s['aum_m'] for s in lst)
        depos = {}
        for s in lst:
            depos.setdefault(s['depositario'], {'count': 0, 'aum_m': 0.0})
            depos[s['depositario']]['count'] += 1
            depos[s['depositario']]['aum_m'] += s['aum_m']
        primary = max(depos, key=lambda k: depos[k]['aum_m'])
        fi = fi_by_gestora.get(g)
        fi_primary = max(fi['depos'], key=lambda k: fi['depos'][k]) if fi and fi['depos'] else None
        inv = depos.get(INVERSIS, {'count': 0, 'aum_m': 0.0})
        gestora_stats.append({
            'gestora': g, 'short': _short(g), 'grupo': lst[0]['grupo'],
            'count': len(lst), 'aum_m': round(aum, 1),
            'share_pct': round(aum / total_aum_m * 100, 2) if total_aum_m else 0,
            'investors': sum(s['investors'] for s in lst),
            'wtd_fee_bps': _bps(_wavg([(s['aum_m'], s['dep_fee_bps'] / 100 if s['dep_fee_bps'] is not None else None) for s in lst])),
            'est_fees_k': round(sum(s['est_dep_fee_k'] or 0 for s in lst), 0),
            'primary_depositario': primary,
            'depositarios': sorted([{'depositario': k, 'short': _short(k), 'count': v['count'],
                                     'aum_m': round(v['aum_m'], 1)} for k, v in depos.items()],
                                   key=lambda x: -x['aum_m']),
            'inversis_count': inv['count'], 'inversis_aum_m': round(inv['aum_m'], 1),
            'is_inversis_client': inv['count'] > 0,
            'is_captive': any(s['is_captive'] for s in lst if s['depositario'] == primary),
            'fi_aum_m': round(fi['aum_m'], 1) if fi else 0.0,
            'fi_primary_depositario': fi_primary,
        })
    gestora_stats.sort(key=lambda x: -x['aum_m'])

    # Prospects: SICAV books held at another depositary, excluding captive set-ups
    targets = [g for g in gestora_stats
               if g['aum_m'] - g['inversis_aum_m'] > 0 and not g['is_captive']
               and g['primary_depositario'] != INVERSIS]
    targets = [{**g, 'addressable_aum_m': round(g['aum_m'] - g['inversis_aum_m'], 1)} for g in targets]
    targets.sort(key=lambda x: -x['addressable_aum_m'])

    # ── Fee & size distributions ──
    fee_hist = []
    edges = list(range(0, 16))
    for lo in edges:
        sel = [s for s in sicavs if s['dep_fee_bps'] is not None and lo <= round(s['dep_fee_bps']) < lo + 1]
        fee_hist.append(_hist_row(f'{lo}', sel))
    sel = [s for s in sicavs if s['dep_fee_bps'] is not None and round(s['dep_fee_bps']) >= 16]
    fee_hist.append(_hist_row('16+', sel))

    size_buckets = []
    for lo, hi, lab in [(0, 10, '<€10M'), (10, 25, '€10–25M'), (25, 50, '€25–50M'),
                        (50, 100, '€50–100M'), (100, 250, '€100–250M'), (250, 1e12, '>€250M')]:
        sel = [s for s in sicavs if lo <= s['aum_m'] < hi]
        size_buckets.append({
            'bucket': lab, 'count': len(sel), 'aum_m': round(sum(s['aum_m'] for s in sel), 1),
            'median_fee_bps': _median([s['dep_fee_bps'] for s in sel]),
            'wtd_fee_bps': _bps(_wavg([(s['aum_m'], s['dep_fee_bps'] / 100 if s['dep_fee_bps'] is not None else None) for s in sel])),
            'inversis_count': sum(1 for s in sel if s['is_inversis']),
        })

    # ── CNMV aggregates (Cuadro 4.1 across vintages; Cuadro 8.2 / 8.4 from latest) ──
    f3 = _cnmv_file3_paths()
    quarterly = {}
    for _, path in f3:
        quarterly.update(_parse_cuadro_41(path))
    quarterly_trend = [{'quarter': q, 'count': int(v.get('count', 0)),
                        'accounts': int(v.get('accounts', 0)),
                        'aum_bn': round(v.get('aum_k', 0) / 1e6, 2)}
                       for q, v in sorted(quarterly.items())]
    sicav_pl = _parse_pl_fees(f3[-1][1], 'Cuadro 8.4') if f3 else {}
    fi_pl = _parse_pl_fees(f3[-1][1], 'Cuadro 8.2') if f3 else {}
    official_fees = []
    for y in sorted(sicav_pl):
        dec_prev = quarterly.get(f'{y - 1}-Q4', {}).get('aum_k')
        dec_cur = quarterly.get(f'{y}-Q4', {}).get('aum_k')
        implied = (sicav_pl[y].get('dep_k', 0) / ((dec_prev + dec_cur) / 2) * 10000
                   if dec_prev and dec_cur else None)
        official_fees.append({
            'year': y,
            'sicav_dep_fees_m': round(sicav_pl[y].get('dep_k', 0) / 1000, 2),
            'sicav_mgmt_fees_m': round(sicav_pl[y].get('mgmt_k', 0) / 1000, 2),
            'fi_dep_fees_m': round(fi_pl.get(y, {}).get('dep_k', 0) / 1000, 1),
            'sicav_implied_dep_bps': round(implied, 2) if implied else None,
        })

    monthly, inverco_ranking, inverco_month = _inverco_monthly()

    inv = next((d for d in depositario_stats if d['is_inversis']), None)
    inv_list = [s for s in sicavs if s['is_inversis']]
    summary = {
        'period': full['period'],
        'fee_fy_label': full['fy_label'],
        'fee_ytd_label': full['ytd_label'],
        'prior_label': full['prior_label'],
        'count': len(sicavs),
        'series': len(recs),
        'aum_bn': round(total_aum_m / 1000, 2),
        'investors': sum(s['investors'] for s in sicavs),
        'avg_aum_m': round(total_aum_m / len(sicavs), 1) if sicavs else 0,
        'median_aum_m': round(_median([s['aum_m'] for s in sicavs]) or 0, 1),
        'wtd_dep_fee_bps': _bps(_wavg([(s['aum_m'], s['dep_fee_bps'] / 100 if s['dep_fee_bps'] is not None else None) for s in sicavs])),
        'median_dep_fee_bps': _median([s['dep_fee_bps'] for s in sicavs]),
        'est_dep_fees_m': round(total_fees_k / 1000, 2),
        'median_mgmt_fee_pct': _median([s['mgmt_fee_pct'] for s in sicavs]),
        'median_ter_pct': _median([s['ter_pct'] for s in sicavs]),
        'fi_wtd_dep_fee_bps': _bps(fi_dep_wtd),
        'fi_median_dep_fee_bps': _bps(fi_dep_median),
        'depositarios': len(depositario_stats),
        'gestoras': len(gestora_stats),
        'switches': len(switches),
        'entries': len(entries),
        'exits': len(exits),
        'inversis_count': inv['count'] if inv else 0,
        'inversis_aum_bn': inv['aum_bn'] if inv else 0,
        'inversis_share_pct': inv['share_pct'] if inv else 0,
        'inversis_count_share_pct': inv['count_share_pct'] if inv else 0,
        'inversis_rank': inv['rank'] if inv else None,
        'inversis_wtd_fee_bps': inv['wtd_fee_bps'] if inv else None,
        'inversis_est_fees_k': inv['est_fees_k'] if inv else 0,
        'inversis_gestoras': inv['gestoras'] if inv else 0,
        'inversis_median_size_m': round(_median([s['aum_m'] for s in inv_list]) or 0, 1),
        'addressable_aum_bn': round(sum(t['addressable_aum_m'] for t in targets) / 1000, 2),
        'addressable_gestoras': len(targets),
        'inverco_month': inverco_month,
        'inverco_aum_bn': monthly[-1]['aum_bn'] if monthly else None,
        'inverco_count': monthly[-1]['count'] if monthly else None,
    }

    return {
        'summary': summary,
        'depositario_stats': depositario_stats,
        'gestora_stats': gestora_stats,
        'targets': targets[:40],
        'sicavs': sicavs,
        'switches': [{'name': s['name'], 'short': s['short'], 'gestora': s['gestora'],
                      'from': s['prior_depositario'], 'to': s['depositario'],
                      'aum_m': s['aum_m'], 'dep_fee_bps': s['dep_fee_bps']} for s in switches],
        'entries': [{'name': s['name'], 'short': s['short'], 'gestora': s['gestora'],
                     'depositario': s['depositario'], 'aum_m': s['aum_m']} for s in entries],
        'exits': exits,
        'fee_histogram': fee_hist,
        'size_buckets': size_buckets,
        'quarterly_trend': quarterly_trend,
        'official_fees': official_fees,
        'monthly_trend': monthly,
        'inverco_ranking': inverco_ranking[:30],
    }


def _hist_row(label, sel):
    return {
        'bps': label,
        'count': len(sel),
        'inversis_count': sum(1 for s in sel if s['is_inversis']),
        'other_count': sum(1 for s in sel if not s['is_inversis']),
        'aum_m': round(sum(s['aum_m'] for s in sel), 1),
    }
